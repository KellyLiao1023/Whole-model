from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
from typing import Callable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SCAN_CSV = ROOT / "MS2_Data_PyTorch" / "tables" / "assembled_nn_scan_clean.csv"
DEFAULT_CANDIDATE_DIR = ROOT / "Promoter_library_design" / "candidates_output"
DEFAULT_VARIANT_DIR = ROOT / "Promoter_library_design" / "variants_output"
DEFAULT_OUTPUT_DIR = ROOT / "outputs" / "019ec8ae-c0ac-7190-868c-5c4ba74a5396"

BG5 = "AGGGAAGAGACC"
BG3 = "GTCGACTCTAGA"
LINKER = "CAC"
M35_LEN = 6
M10_LEN = 6
SPACERS = [16, 17, 18]
SPACER_BY_CHANNEL = {ch: sp for ch, sp in enumerate(SPACERS)}
CHANNEL_BY_SPACER = {sp: ch for ch, sp in SPACER_BY_CHANNEL.items()}
ELEMENT_COLS = ["UP", "m35", "spacer", "m10", "DIS", "ITS"]

HEADER_FILL = "1F4E79"
SUBHEADER_FILL = "D9EAF7"


@dataclass
class ReplacementResult:
    profile: pd.DataFrame
    suggestions: pd.DataFrame
    detail: pd.DataFrame
    rejected_candidates: pd.DataFrame
    output_path: Path


def _normalize_element_type(element_type: str) -> str:
    key = str(element_type).strip().lower().replace("-", "").replace("_", "")
    mapping = {
        "up": "UP",
        "ul": "UP",
        "m35": "m35",
        "35": "m35",
        "minus35": "m35",
        "pl35": "m35",
        "spacer": "spacer",
        "sp": "spacer",
        "sp16": "spacer",
        "sp17": "spacer",
        "sp18": "spacer",
        "m10": "m10",
        "10": "m10",
        "minus10": "m10",
        "pl10": "m10",
        "dis": "DIS",
        "dl": "DIS",
        "its": "ITS",
    }
    if key not in mapping:
        raise ValueError(f"Unknown element_type: {element_type!r}. Use one of {ELEMENT_COLS}.")
    return mapping[key]


def _safe_slice(seq: str, start, end) -> str:
    if pd.isna(start) or pd.isna(end):
        return ""
    start = max(0, int(start))
    end = min(len(seq), int(end))
    return seq[start:end] if start < end else ""


def _mode_text(series: pd.Series) -> str:
    counts = series.dropna().astype(str).value_counts()
    return "" if counts.empty else str(counts.index[0])


def _mode_n(series: pd.Series) -> int:
    counts = series.dropna().astype(str).value_counts()
    return 0 if counts.empty else int(counts.iloc[0])


def _landing_group(df: pd.DataFrame) -> pd.Series:
    return (
        "m35@"
        + df["observed_m35_start"].astype("Int64").astype(str)
        + "_m10@"
        + df["observed_m10_start"].astype("Int64").astype(str)
        + "_sp"
        + df["observed_spacer_len"].astype("Int64").astype(str)
    )


def _coerce_scan_df(scan_df: pd.DataFrame) -> pd.DataFrame:
    df = scan_df.copy()
    numeric_cols = [
        "design_m35_start",
        "design_spacer_len",
        "design_m10_start",
        "design_channel",
        "model_m35_offset",
        "design_arch_start",
        "observed_arch_start",
        "observed_channel",
        "observed_spacer_len",
        "arch_shift",
        "channel_shift",
        "observed_m35_start",
        "observed_m10_start",
        "m35_shift",
        "m10_shift",
        "best_model_energy",
        "best_model_log10",
        "design_model_energy",
        "design_model_log10",
        "delta_model_log10",
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["pull_m35"] = df["m35_shift"].fillna(0) != 0
    df["pull_m10"] = df["m10_shift"].fillna(0) != 0
    df["pull_either"] = df["pull_m35"] | df["pull_m10"]
    df["pull_both"] = df["pull_m35"] & df["pull_m10"]
    df["shifted_m35_seq"] = df.apply(
        lambda r: _safe_slice(r["Sequence"], r["observed_m35_start"], r["observed_m35_start"] + M35_LEN),
        axis=1,
    )
    df["shifted_m10_seq"] = df.apply(
        lambda r: _safe_slice(r["Sequence"], r["observed_m10_start"], r["observed_m10_start"] + M10_LEN),
        axis=1,
    )
    df["landing_group"] = _landing_group(df)
    return df


def load_scan_df(path: str | Path = DEFAULT_SCAN_CSV) -> pd.DataFrame:
    return _coerce_scan_df(pd.read_csv(path))


def _candidate_file_names(element_type: str, seq_len: int | None) -> list[tuple[str, str]]:
    if element_type == "UP":
        keys = ["UP"]
    elif element_type == "m35":
        keys = ["minus35"]
    elif element_type == "m10":
        keys = ["minus10"]
    elif element_type == "DIS":
        keys = ["Dis"]
    elif element_type == "ITS":
        keys = ["ITS"]
    elif element_type == "spacer":
        spacer_keys = [f"Sp{seq_len}"] if seq_len else [f"Sp{x}" for x in SPACERS]
        keys = spacer_keys
    else:
        keys = [element_type]

    out = []
    for key in keys:
        out.append(("candidates", f"candidates_{key}.csv"))
        out.append(("variants", f"variants_{key}.csv"))
    return out


def _read_candidate_file(path: Path, source_type: str) -> pd.DataFrame:
    df = pd.read_csv(path)
    if "sequence" not in df.columns:
        return pd.DataFrame()
    df = df.copy()
    df["candidate_source_type"] = source_type
    df["candidate_source_file"] = path.name
    return df


def _load_candidate_pool(
    element_type: str,
    bad_seq: str,
    scan_df: pd.DataFrame,
    candidate_dir: Path,
    variant_dir: Path,
) -> pd.DataFrame:
    seq_len = len(bad_seq)
    frames = []
    for source_type, file_name in _candidate_file_names(element_type, seq_len):
        base = candidate_dir if source_type == "candidates" else variant_dir
        path = base / file_name
        if path.exists():
            frame = _read_candidate_file(path, source_type)
            if not frame.empty:
                frames.append(frame)

    if frames:
        pool = pd.concat(frames, ignore_index=True, sort=False)
    else:
        current = sorted(scan_df[element_type].dropna().astype(str).unique())
        pool = pd.DataFrame(
            {
                "element": element_type,
                "sequence": current,
                "candidate_source_type": "current_design",
                "candidate_source_file": "assembled_nn_scan_clean.csv",
            }
        )

    pool["sequence"] = pool["sequence"].astype(str).str.upper()
    pool = pool.drop_duplicates("sequence").reset_index(drop=True)
    pool["length"] = pool["sequence"].str.len()
    return pool


def _add_historical_candidate_metrics(pool: pd.DataFrame, scan_df: pd.DataFrame, element_type: str) -> pd.DataFrame:
    rows = []
    for seq in pool["sequence"].astype(str):
        mask = scan_df[element_type].astype(str).str.upper() == seq.upper()
        sub = scan_df[mask]
        rows.append(
            {
                "candidate_seq": seq,
                "historical_n": int(len(sub)),
                "historical_m35_shift_rate": float(sub["pull_m35"].mean()) if len(sub) else np.nan,
                "historical_m10_shift_rate": float(sub["pull_m10"].mean()) if len(sub) else np.nan,
                "historical_any_shift_rate": float(sub["pull_either"].mean()) if len(sub) else np.nan,
                "historical_top_landing_group": _mode_text(sub["landing_group"]) if len(sub) else "",
            }
        )
    hist = pd.DataFrame(rows)
    return pool.merge(hist, left_on="sequence", right_on="candidate_seq", how="left").drop(columns=["candidate_seq"])


def _select_candidate_pool(
    pool: pd.DataFrame,
    bad_seq: str,
    max_candidates: int,
    preserve_bin: bool,
    reject_known_risky: bool,
    known_risky_threshold: float,
    min_known_contexts: int,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rejected = []
    bad_meta = pool.loc[pool["sequence"].astype(str).str.upper() == bad_seq.upper()]
    bad_bin = bad_meta["bin_id"].dropna().iloc[0] if preserve_bin and "bin_id" in bad_meta and not bad_meta.empty else np.nan

    selected = pool.copy()
    exact = selected["sequence"].astype(str).str.upper() == bad_seq.upper()
    if exact.any():
        bad = selected[exact].copy()
        bad["reject_reason"] = "locked_sequence_itself"
        rejected.append(bad)
        selected = selected[~exact].copy()

    wrong_len = selected["length"] != len(bad_seq)
    if wrong_len.any():
        bad = selected[wrong_len].copy()
        bad["reject_reason"] = "wrong_length"
        rejected.append(bad)
        selected = selected[~wrong_len].copy()

    if reject_known_risky and "historical_any_shift_rate" in selected.columns:
        risky = (
            (selected["historical_n"].fillna(0) >= min_known_contexts)
            & (selected["historical_any_shift_rate"].fillna(0) >= known_risky_threshold)
        )
        if risky.any():
            bad = selected[risky].copy()
            bad["reject_reason"] = "known_high_shift_rate"
            rejected.append(bad)
            selected = selected[~risky].copy()

    if selected.empty:
        rejected_df = pd.concat(rejected, ignore_index=True, sort=False) if rejected else pd.DataFrame()
        return selected, rejected_df

    if preserve_bin and pd.notna(bad_bin) and "bin_id" in selected.columns:
        selected["bin_distance"] = (pd.to_numeric(selected["bin_id"], errors="coerce") - float(bad_bin)).abs()
        selected["same_bin"] = selected["bin_distance"] == 0
    else:
        selected["bin_distance"] = np.nan
        selected["same_bin"] = False

    if "rank_in_bin" not in selected.columns:
        selected["rank_in_bin"] = np.nan
    if "read_count" not in selected.columns:
        selected["read_count"] = np.nan
    if "delta_epsilon" not in selected.columns:
        selected["delta_epsilon"] = np.nan

    selected = selected.sort_values(
        [
            "same_bin",
            "bin_distance",
            "candidate_source_type",
            "rank_in_bin",
            "read_count",
            "sequence",
        ],
        ascending=[False, True, True, True, False, True],
        na_position="last",
    ).head(max_candidates)

    rejected_df = pd.concat(rejected, ignore_index=True, sort=False) if rejected else pd.DataFrame()
    return selected.reset_index(drop=True), rejected_df.reset_index(drop=True)


def _assemble_sequence(row: pd.Series) -> str:
    return (
        BG5
        + str(row["UP"])
        + LINKER
        + str(row["m35"])
        + str(row["spacer"])
        + str(row["m10"])
        + str(row["DIS"])
        + str(row["ITS"])
        + BG3
    )


def _refresh_design_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["Sequence"] = out.apply(_assemble_sequence, axis=1)
    out["design_m35_start"] = out["UP"].astype(str).str.len() + len(BG5) + len(LINKER)
    out["design_spacer_len"] = out["spacer"].astype(str).str.len()
    out["design_m10_start"] = out["design_m35_start"] + M35_LEN + out["design_spacer_len"]
    out["design_channel"] = out["design_spacer_len"].map(CHANNEL_BY_SPACER)
    if out["design_channel"].isna().any():
        bad_lengths = sorted(out.loc[out["design_channel"].isna(), "design_spacer_len"].unique())
        raise ValueError(f"Unsupported spacer length after replacement: {bad_lengths}")
    out["design_channel"] = out["design_channel"].astype(int)
    out["design_arch_start"] = out["design_m35_start"] - out["model_m35_offset"]
    return out


def _profile_locked_sequence(scan_df: pd.DataFrame, element_type: str, bad_seq: str) -> pd.DataFrame:
    locked = scan_df[scan_df[element_type].astype(str).str.upper() == bad_seq.upper()].copy()
    if locked.empty:
        raise ValueError(f"{element_type} sequence {bad_seq!r} was not found in scan_df.")

    total = len(locked)
    rows = [
        ("element_type", element_type),
        ("locked_sequence", bad_seq),
        ("contexts_found", total),
        ("pull_m35_n", int(locked["pull_m35"].sum())),
        ("pull_m35_rate", float(locked["pull_m35"].mean())),
        ("pull_m10_n", int(locked["pull_m10"].sum())),
        ("pull_m10_rate", float(locked["pull_m10"].mean())),
        ("pull_both_n", int(locked["pull_both"].sum())),
        ("pull_both_rate", float(locked["pull_both"].mean())),
        ("pull_either_n", int(locked["pull_either"].sum())),
        ("pull_either_rate", float(locked["pull_either"].mean())),
        ("top_landing_group", _mode_text(locked.loc[locked["pull_either"], "landing_group"])),
        ("top_landing_group_n", _mode_n(locked.loc[locked["pull_either"], "landing_group"])),
        ("top_shifted_m35_seq", _mode_text(locked.loc[locked["pull_m35"], "shifted_m35_seq"])),
        ("top_shifted_m35_n", _mode_n(locked.loc[locked["pull_m35"], "shifted_m35_seq"])),
        ("top_shifted_m10_seq", _mode_text(locked.loc[locked["pull_m10"], "shifted_m10_seq"])),
        ("top_shifted_m10_n", _mode_n(locked.loc[locked["pull_m10"], "shifted_m10_seq"])),
    ]
    return pd.DataFrame(rows, columns=["metric", "value"])


def _prepare_contexts(scan_df: pd.DataFrame, element_type: str, bad_seq: str) -> pd.DataFrame:
    contexts = scan_df[scan_df[element_type].astype(str).str.upper() == bad_seq.upper()].copy()
    if contexts.empty:
        raise ValueError(f"{element_type} sequence {bad_seq!r} was not found in scan_df.")

    contexts = contexts.reset_index().rename(columns={"index": "original_row_id"})
    contexts["original_pull_m35"] = contexts["pull_m35"]
    contexts["original_pull_m10"] = contexts["pull_m10"]
    contexts["original_pull_either"] = contexts["pull_either"]
    contexts["original_observed_m35_start"] = contexts["observed_m35_start"]
    contexts["original_observed_m10_start"] = contexts["observed_m10_start"]
    contexts["original_landing_group"] = contexts["landing_group"]
    return contexts


def _scan_one_candidate(
    contexts: pd.DataFrame,
    element_type: str,
    candidate_seq: str,
    scanner: Callable[[pd.DataFrame], pd.DataFrame],
) -> pd.DataFrame:
    sub = contexts.copy()
    sub[element_type] = candidate_seq
    sub = _refresh_design_columns(sub)
    scanned = scanner(sub)
    scanned = _coerce_scan_df(scanned)

    keep_original = [
        "original_row_id",
        "original_pull_m35",
        "original_pull_m10",
        "original_pull_either",
        "original_observed_m35_start",
        "original_observed_m10_start",
        "original_landing_group",
    ]
    for col in keep_original:
        scanned[col] = sub[col].values
    scanned["candidate_seq"] = candidate_seq
    scanned["new_pull_m35"] = scanned["m35_shift"] != 0
    scanned["new_pull_m10"] = scanned["m10_shift"] != 0
    scanned["new_pull_either"] = scanned["new_pull_m35"] | scanned["new_pull_m10"]
    scanned["same_landing_as_original"] = (
        (scanned["observed_m35_start"] == scanned["original_observed_m35_start"])
        & (scanned["observed_m10_start"] == scanned["original_observed_m10_start"])
    )
    m35_rows = scanned["original_pull_m35"]
    m10_rows = scanned["original_pull_m10"]
    target_rescued = pd.Series(pd.NA, index=scanned.index, dtype="object")
    target_rescued.loc[scanned["original_pull_either"]] = True
    target_rescued.loc[m35_rows] = target_rescued.loc[m35_rows] & ~scanned.loc[m35_rows, "new_pull_m35"]
    target_rescued.loc[m10_rows] = target_rescued.loc[m10_rows] & ~scanned.loc[m10_rows, "new_pull_m10"]
    scanned["target_rescued"] = target_rescued
    return scanned


def _summarize_candidate(scanned: pd.DataFrame, candidate_meta: pd.Series) -> dict:
    original_pulled = scanned["original_pull_either"].fillna(False)
    pulled_n = int(original_pulled.sum())
    target_rescued = scanned.loc[original_pulled, "target_rescued"].map(lambda x: bool(x) if pd.notna(x) else False)
    target_rescued_n = int(target_rescued.sum())

    out = {
        "candidate_seq": scanned["candidate_seq"].iloc[0],
        "candidate_source_type": candidate_meta.get("candidate_source_type", ""),
        "candidate_source_file": candidate_meta.get("candidate_source_file", ""),
        "bin_id": candidate_meta.get("bin_id", np.nan),
        "same_bin": candidate_meta.get("same_bin", pd.NA),
        "bin_distance": candidate_meta.get("bin_distance", np.nan),
        "rank_in_bin": candidate_meta.get("rank_in_bin", np.nan),
        "delta_epsilon": candidate_meta.get("delta_epsilon", np.nan),
        "LogGFP": candidate_meta.get("LogGFP", candidate_meta.get("LogGFP_measured", np.nan)),
        "read_count": candidate_meta.get("read_count", np.nan),
        "historical_n": candidate_meta.get("historical_n", np.nan),
        "historical_any_shift_rate": candidate_meta.get("historical_any_shift_rate", np.nan),
        "contexts_tested": int(len(scanned)),
        "original_pulled_contexts": pulled_n,
        "target_rescued_n": target_rescued_n,
        "target_rescue_rate": target_rescued_n / pulled_n if pulled_n else np.nan,
        "clean_after_n": int((~scanned["new_pull_either"]).sum()),
        "clean_after_rate": float((~scanned["new_pull_either"]).mean()),
        "new_any_shift_rate": float(scanned["new_pull_either"].mean()),
        "new_m35_shift_rate": float(scanned["new_pull_m35"].mean()),
        "new_m10_shift_rate": float(scanned["new_pull_m10"].mean()),
        "same_landing_rate": float(scanned.loc[original_pulled, "same_landing_as_original"].mean()) if pulled_n else np.nan,
        "median_delta_model_log10_after": float(scanned["delta_model_log10"].median()),
        "mean_delta_model_log10_after": float(scanned["delta_model_log10"].mean()),
        "top_after_landing_group": _mode_text(scanned["landing_group"]),
        "top_after_landing_group_n": _mode_n(scanned["landing_group"]),
        "top_after_shifted_m35_seq": _mode_text(scanned.loc[scanned["new_pull_m35"], "shifted_m35_seq"]),
        "top_after_shifted_m35_n": _mode_n(scanned.loc[scanned["new_pull_m35"], "shifted_m35_seq"]),
        "top_after_shifted_m10_seq": _mode_text(scanned.loc[scanned["new_pull_m10"], "shifted_m10_seq"]),
        "top_after_shifted_m10_n": _mode_n(scanned.loc[scanned["new_pull_m10"], "shifted_m10_seq"]),
    }
    return out


def _sanitize_filename(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("_") or "sequence"


def _write_excel(
    profile: pd.DataFrame,
    suggestions: pd.DataFrame,
    detail: pd.DataFrame,
    rejected: pd.DataFrame,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        profile.to_excel(writer, sheet_name="Locked_Seq_Profile", index=False)
        suggestions.to_excel(writer, sheet_name="Replacement_Suggestions", index=False)
        detail.to_excel(writer, sheet_name="Context_Rescue_Detail", index=False)
        rejected.to_excel(writer, sheet_name="Rejected_Candidates", index=False)

    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row >= 1 and max_col >= 1:
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_idx in range(1, max_col + 1):
                col_letter = get_column_letter(col_idx)
                values = [ws.cell(row=r, column=col_idx).value for r in range(1, min(max_row, 200) + 1)]
                width = min(max(max(len(str(v)) if v is not None else 0 for v in values) + 2, 10), 45)
                ws.column_dimensions[col_letter].width = width
            if max_row > 1:
                ref = f"A1:{get_column_letter(max_col)}{max_row}"
                safe_name = re.sub(r"[^A-Za-z0-9_]", "_", ws.title)[:24]
                table = Table(displayName=f"T_{safe_name}", ref=ref)
                style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
                table.tableStyleInfo = style
                ws.add_table(table)

        headers = {ws.cell(row=1, column=c).value: c for c in range(1, max_col + 1)}
        good_when_high = {"target_rescue_rate", "clean_after_rate"}
        good_when_low = {"new_any_shift_rate", "same_landing_rate"}
        for name in [*good_when_high, *good_when_low]:
            if name in headers and max_row > 2:
                col = get_column_letter(headers[name])
                start_color, end_color = ("F8696B", "63BE7B") if name in good_when_high else ("63BE7B", "F8696B")
                ws.conditional_formatting.add(
                    f"{col}2:{col}{max_row}",
                    ColorScaleRule(
                        start_type="min",
                        start_color=start_color,
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color=end_color,
                    ),
                )

    wb.save(output_path)


def suggest_replacements(
    scan_df: pd.DataFrame | None,
    element_type: str,
    bad_seq: str,
    scanner: Callable[[pd.DataFrame], pd.DataFrame],
    *,
    top_n: int = 20,
    max_candidates: int = 80,
    preserve_bin: bool = True,
    reject_known_risky: bool = True,
    known_risky_threshold: float = 0.80,
    min_known_contexts: int = 10,
    detail_top_n: int = 10,
    candidate_dir: str | Path = DEFAULT_CANDIDATE_DIR,
    variant_dir: str | Path = DEFAULT_VARIANT_DIR,
    output_dir: str | Path = DEFAULT_OUTPUT_DIR,
    output_path: str | Path | None = None,
) -> ReplacementResult:
    """Suggest replacement sequences by rescanning substitutions in all locked contexts.

    Pass the notebook's scan_corepromoter_model as scanner so this helper uses the
    same model, offsets, and score conversion as the active analysis notebook.
    """
    if scanner is None:
        raise ValueError("scanner is required. Pass scanner=scan_corepromoter_model from the notebook.")

    element_type = _normalize_element_type(element_type)
    bad_seq = str(bad_seq).upper()
    scan_df = load_scan_df() if scan_df is None else _coerce_scan_df(scan_df)
    profile = _profile_locked_sequence(scan_df, element_type, bad_seq)
    contexts = _prepare_contexts(scan_df, element_type, bad_seq)

    pool = _load_candidate_pool(
        element_type,
        bad_seq,
        scan_df,
        Path(candidate_dir),
        Path(variant_dir),
    )
    pool = _add_historical_candidate_metrics(pool, scan_df, element_type)
    pool, rejected = _select_candidate_pool(
        pool,
        bad_seq,
        max_candidates=max_candidates,
        preserve_bin=preserve_bin,
        reject_known_risky=reject_known_risky,
        known_risky_threshold=known_risky_threshold,
        min_known_contexts=min_known_contexts,
    )

    summaries = []
    details = []
    for _, candidate_meta in pool.iterrows():
        candidate_seq = str(candidate_meta["sequence"]).upper()
        scanned = _scan_one_candidate(contexts, element_type, candidate_seq, scanner)
        summaries.append(_summarize_candidate(scanned, candidate_meta))
        details.append(scanned)

    suggestions = pd.DataFrame(summaries)
    if suggestions.empty:
        detail = pd.DataFrame()
    else:
        suggestions = suggestions.sort_values(
            [
                "target_rescue_rate",
                "new_any_shift_rate",
                "same_landing_rate",
                "clean_after_rate",
                "historical_any_shift_rate",
            ],
            ascending=[False, True, True, False, True],
            na_position="last",
        ).reset_index(drop=True)
        suggestions.insert(0, "recommended_rank", range(1, len(suggestions) + 1))
        suggestions = suggestions.head(top_n)

        detail_all = pd.concat(details, ignore_index=True, sort=False)
        keep_candidates = suggestions.head(detail_top_n)["candidate_seq"].tolist()
        detail_cols = [
            "candidate_seq",
            "original_row_id",
            "original_pull_m35",
            "original_pull_m10",
            "original_landing_group",
            "m35_shift",
            "m10_shift",
            "new_pull_m35",
            "new_pull_m10",
            "target_rescued",
            "same_landing_as_original",
            "observed_m35_start",
            "observed_m10_start",
            "landing_group",
            "shifted_m35_seq",
            "shifted_m10_seq",
            "design_model_log10",
            "best_model_log10",
            "delta_model_log10",
            "UP",
            "m35",
            "spacer",
            "m10",
            "DIS",
            "ITS",
            "Sequence",
        ]
        detail = detail_all[detail_all["candidate_seq"].isin(keep_candidates)].copy()
        detail = detail[[c for c in detail_cols if c in detail.columns]]

    if output_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_name = f"replacement_{element_type}_{_sanitize_filename(bad_seq)}_{stamp}.xlsx"
        output_path = Path(output_dir) / file_name
    else:
        output_path = Path(output_path)

    _write_excel(profile, suggestions, detail, rejected, output_path)
    return ReplacementResult(
        profile=profile,
        suggestions=suggestions,
        detail=detail,
        rejected_candidates=rejected,
        output_path=output_path,
    )
